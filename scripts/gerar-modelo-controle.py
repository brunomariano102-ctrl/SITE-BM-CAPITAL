"""Gera somente um modelo vazio; requer openpyxl 3.1.5, fora das dependências do site."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

DEST = Path(__file__).resolve().parents[1] / 'docs/juridico/MODELO-CONTROLE-MARKETING.xlsx'
NAVY, GOLD = '071827', 'CBA45E'
PREFERENCES = ['NAO_VERIFICADO', 'AUTORIZOU', 'RECUSOU', 'REVOGOU']
COMMUNICATION = ['PENDENTE', 'COMUNICADO', 'CONFIRMADO_PELA_EQUIPE', 'NAO_SE_APLICA']
MAX_ROW = 301


def dropdown(sheet, column, values, title, message):
    rule = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    rule.errorTitle = 'Escolha uma opção da lista'
    rule.error = 'Use um dos valores disponíveis; não presuma autorização.'
    rule.promptTitle = title
    rule.prompt = message
    rule.showInputMessage = True
    rule.showErrorMessage = True
    rule.errorStyle = 'stop'
    sheet.add_data_validation(rule)
    rule.add(f'{column}2:{column}{MAX_ROW}')


def grid(workbook, name, headers, table_name, widths):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    sheet.freeze_panes = 'A2'
    sheet.row_dimensions[1].height = 48
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Calibri', color=GOLD, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row in sheet.iter_rows(min_row=2, max_row=MAX_ROW, max_col=len(headers)):
        for cell in row:
            cell.number_format = '@'
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    last = sheet.cell(1, len(headers)).column_letter
    table = Table(displayName=table_name, ref=f'A1:{last}{MAX_ROW}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    return sheet


def main():
    wb = Workbook()
    intro = wb.active
    intro.title = 'LEIA-ME'
    rows = [
        ('BM CAPITAL — MODELO MANUAL', 'PROPOSTA PARA AVALIAÇÃO; SEM CONTATOS REAIS'),
        ('Finalidade', 'Organizar escolhas e atribuições. Não integra o site, CRM, WhatsApp ou e-mail.'),
        ('Responsável', 'O responsável pela BM Capital declarou que coordenará privacidade e revogação.'),
        ('Antes de usar', 'Definir armazenamento, pessoas autorizadas, retenção, descarte e procedimento com o responsável/jurídico.'),
        ('CONTATOS', 'Uma linha por contato atribuído. Use identificador interno e referência no canal oficial; minimize dados duplicados.'),
        ('Preferência', 'NAO_VERIFICADO, AUTORIZOU, RECUSOU ou REVOGOU. Em branco equivale a preferência não verificada para este controle.'),
        ('AUTORIZOU', 'Só selecionar com evidência efetivamente recebida e pertinente à finalidade. A planilha não confere essa evidência.'),
        ('Atendimento', 'Recusar marketing não impede responder ao pedido de atendimento. Separar a finalidade de futuras campanhas.'),
        ('ALTERACOES', 'Registrar manualmente mudanças e respectivas evidências; não sobrescrever o único registro histórico disponível.'),
        ('Datas', 'Digite data, hora e fuso reais. Formato sugerido: AAAA-MM-DD HH:MM UTC-03:00. Nada é preenchido automaticamente.'),
        ('Evidência', 'Use referência ao registro necessário; não copie documentos, senhas ou conversas completas sem necessidade.'),
        ('Revogação', 'Atualizar preferência e registro de alteração e comunicar vendedores envolvidos antes de novas ações promocionais.'),
        ('Números próprios', 'A planilha não bloqueia envios dos vendedores. A execução e confirmação precisam ser organizadas pela empresa.'),
        ('Acesso', 'Filtros, cores, abas e validações não controlam acesso. Não distribuir a base completa ou criar link público.'),
        ('Limites', 'Sem macros, fórmulas, integrações, registro imutável ou bloqueio automático. Não certifica conformidade jurídica.'),
        ('Texto externo', 'Mantenha campos como texto; ao colar, preserve esse formato. Validações podem ser contornadas por colagem.'),
        ('Capacidade', '300 linhas vazias preparadas em cada aba. Ao expandir, confira tabelas, filtros, formato e validações.'),
        ('Modelo, não política', 'Nenhum prazo de retenção, base legal definitiva ou nova autorização foi definido pelo arquivo.'),
    ]
    for row in rows:
        intro.append(row)
    intro.column_dimensions['A'].width = 27
    intro.column_dimensions['B'].width = 108
    intro.freeze_panes = 'B2'
    for row in intro:
        intro.row_dimensions[row[0].row].height = 43
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            cell.font = Font(name='Calibri', size=11)
    for cell in intro[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Calibri', bold=True, color=GOLD, size=12)
    contacts = grid(wb, 'CONTATOS', [
        'ID interno', 'Nome preferido', 'Referência no canal oficial', 'Vendedor responsável',
        'Preferência de marketing', 'Data/hora e fuso do registro', 'Versão do aviso',
        'Referência da evidência', 'Data/hora e fuso da revogação', 'Comunicação à equipe',
    ], 'ContatosBM', dict(zip('ABCDEFGHIJ', [19, 24, 35, 26, 25, 32, 22, 36, 32, 31])))
    dropdown(contacts, 'E', PREFERENCES, 'Preferência efetivamente recebida', 'Sem evidência, use NAO_VERIFICADO. AUTORIZOU não é padrão.')
    dropdown(contacts, 'J', COMMUNICATION, 'Comunicação manual', 'Só marque confirmação depois de recebê-la da equipe envolvida.')
    history = grid(wb, 'ALTERACOES', [
        'ID interno do contato', 'Data/hora e fuso', 'Evento', 'Registrado por', 'Vendedor envolvido',
        'Preferência anterior', 'Nova preferência', 'Referência da evidência', 'Comunicação à equipe',
    ], 'AlteracoesBM', dict(zip('ABCDEFGHI', [24, 32, 26, 26, 26, 25, 25, 36, 31])))
    dropdown(history, 'C', ['RECEBIMENTO', 'AUTORIZACAO', 'RECUSA', 'REVOGACAO', 'ATRIBUICAO', 'ATUALIZACAO_EQUIPE'], 'Evento real', 'O modelo não gera eventos; registre somente o que ocorreu.')
    for column in ['F', 'G']:
        dropdown(history, column, PREFERENCES, 'Preferência', 'Registre a escolha anterior e a nova sem deduzi-las.')
    dropdown(history, 'I', COMMUNICATION, 'Comunicação manual', 'Marque a situação real, sem presumir cumprimento.')
    wb.properties.title = 'BM Capital — modelo manual de atendimento e marketing'
    wb.properties.creator = 'BM Capital'
    wb.properties.description = 'Modelo vazio, sem macros, sem integração e sem certificação de conformidade.'
    DEST.parent.mkdir(parents=True, exist_ok=True)
    if DEST.exists():
        raise SystemExit('O arquivo já existe. Preserve eventual preenchimento antes de gerar outra cópia.')
    wb.save(DEST)
    print(DEST)


if __name__ == '__main__':
    main()
